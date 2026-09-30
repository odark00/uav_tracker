#!/usr/bin/env python3
"""
Детекція без трекера на Anti-UAV валідації + свіп параметрів.

Структура:  val/<seq>/visible.mp4 + visible.json (exist, gt_rect [x,y,w,h])
            val/<seq>/infrared.mp4 + infrared.json (те саме)

Запуск:  python eval_detect.py --ai 1 --conf 0.25 --iou 0.45 --imgsz 640
          python eval_detect.py --ai 1,2 --conf 0.1 --stride 1 --modality both

Логіка: тільки детектор (БЕЗ трекера, чистий YOLO predict).
Всі бокси моделі = UAV (class-agnostic). Кадр з exist=1: max-IoU>=0.5 -> TP,
інакше FN (+ зайві бокси в FP). Кадр з exist=0: будь-який бокс -> FP.
Метрики: P/R/F1@0.5, середній ms/кадр, FPS. Кожен запуск дописує рядок у xlsx.
"""
import argparse
import time
from pathlib import Path

import cv2
import numpy as np
import torch
from ultralytics import YOLO

from infer import MODELS

ROOT = Path(__file__).resolve().parent

# Колонки xlsx: ключ -> заголовок з розшифровкою (пишеться в 1-й рядок при створенні файлу)
COLUMNS = [
    ("model", "AI (номер моделі 1-7)"),
    ("weights", "Ваги (файл з папки ai/)"),
    ("modality", "Модальність (visible/infrared/both)"),
    ("conf", "conf (поріг впевненості)"),
    ("iou", "iou (NMS-поріг)"),
    ("imgsz", "imgsz (розмір входу)"),
    ("seqs", "Відео (к-сть)"),
    ("frames", "Кадрів (оброблено)"),
    ("gt", "GT (реальних об'єктів)"),
    ("tp", "TP (вірно знайдено)"),
    ("fp", "FP (хибні спрацювання)"),
    ("fn", "FN (пропущено)"),
    ("p", "Precision (точність)"),
    ("r", "Recall (повнота)"),
    ("f1", "F1 (гарм. середнє)"),
    ("ap", "AP (середня точність)"),
    ("map50", "mAP@50"),
    ("map5095", "mAP@50:95"),
    ("params", "Params (M)"),
    ("size", "Model size (MB)"),
    ("device", "GPU/CPU"),
    ("vram", "VRAM (GB)"),
    ("ms", "мс/кадр (середній час)"),
    ("fps", "FPS (кадрів/с)"),
]

# Поріг для збору детекцій під AP/mAP (P/R/F1 при цьому рахуються на --conf)
AP_CONF = 0.001


def parse_ai(s):
    if s is None or s == "all":
        return sorted(MODELS)
    return [int(x) for x in s.split(",") if int(x) in MODELS]


def iou_xywh(pred_xyxy, gt_xywh):
    gx, gy, gw, gh = gt_xywh
    g = np.array([gx, gy, gx + gw, gy + gh], float)
    x1, y1 = max(pred_xyxy[0], g[0]), max(pred_xyxy[1], g[1])
    x2, y2 = min(pred_xyxy[2], g[2]), min(pred_xyxy[3], g[3])
    inter = max(0, x2 - x1) * max(0, y2 - y1)
    pa = max(0, pred_xyxy[2] - pred_xyxy[0]) * max(0, pred_xyxy[3] - pred_xyxy[1])
    ga = gw * gh
    return inter / (pa + ga - inter + 1e-9)


def load_gt(seq_dir, modality):
    """-> (video_path, exist, gt_rect) або None якщо пари немає."""
    import json
    v = seq_dir / f"{modality}.mp4"
    j = seq_dir / f"{modality}.json"
    if not v.exists() or not j.exists():
        return None
    d = json.loads(j.read_text())
    return v, d["exist"], d["gt_rect"]


def ap_voc(dets, n_gt):
    """VOC-2010 AP: dets = [(conf, is_tp)], сортуються за спаданням conf."""
    if not dets or n_gt == 0:
        return 0.0
    dets = sorted(dets, key=lambda x: -x[0])
    tp = np.cumsum([t for _, t in dets])
    fp = np.cumsum([1 - t for _, t in dets])
    rec = tp / n_gt
    prec = tp / (tp + fp + 1e-9)
    mrec = np.r_[0, rec, 1]
    mpre = np.r_[0, prec, 0]
    for i in range(len(mpre) - 2, -1, -1):
        mpre[i] = max(mpre[i], mpre[i + 1])
    return float(np.sum((mrec[1:] - mrec[:-1]) * mpre[1:]))


def model_info(mp, model):
    """Params (M), розмір ваг (MB), VRAM (GB, якщо CUDA)."""
    try:
        params = sum(p.numel() for p in model.model.parameters()) / 1e6
    except Exception:
        params = 0.0
    size = mp.stat().st_size / (1024 * 1024)
    vram = 0.0
    if torch.cuda.is_available():
        try:
            torch.cuda.synchronize()
            vram = torch.cuda.max_memory_allocated() / (1024 ** 3)
            torch.cuda.reset_peak_memory_stats()
        except Exception:
            pass
    return params, size, vram


def eval_combo(model, seqs, conf, iou, imgsz, device, stride, tag=""):
    tp = fp = fn = frames = 0
    infer_t = 0.0
    dets = {}  # поріг IoU -> [(conf, is_tp)] для AP/mAP
    n_gt = 0
    total = sum((len(exist) + stride - 1) // stride for _, exist, _ in seqs)
    done = 0
    last_pct = -1
    for vpath, exist, gt in seqs:
        cap = cv2.VideoCapture(str(vpath))
        if not cap.isOpened():
            print(f"  [skip] cannot open {vpath}")
            continue
        nv = int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) or 0
        if nv and abs(nv - len(exist)) > max(5, 0.05 * len(exist)):
            print(f"  [warn] {vpath.name}: відео {nv} кадрів, GT {len(exist)} — невідповідність!")
        for i in range(0, len(exist), stride):
            cap.set(cv2.CAP_PROP_POS_FRAMES, i)
            ok, frame = cap.read()
            if not ok:
                break
            t0 = time.perf_counter()
            res = model.predict(frame, conf=conf, iou=iou, imgsz=imgsz,
                                device=device, verbose=False)[0]
            infer_t += time.perf_counter() - t0
            frames += 1
            boxes = res.boxes.xyxy.cpu().numpy() if res.boxes is not None else np.empty((0, 4))
            if exist[i]:
                n_gt += 1
                best = max((iou_xywh(b, gt[i]) for b in boxes), default=0.0)
                if best >= 0.5:
                    tp += 1
                    fp += max(0, len(boxes) - 1)
                else:
                    fn += 1
                    fp += len(boxes)
                # mAP: окремий прогін з низьким порогом, щоб зібрати повну P-R криву.
                # greedy-матчинг: бокси за спаданням conf, перший з IoU>=t забирає GT.
                r_all = model.predict(frame, conf=AP_CONF, iou=iou, imgsz=imgsz,
                                      device=device, verbose=False)[0]
                if r_all.boxes is not None and len(r_all.boxes):
                    ab = r_all.boxes.xyxy.cpu().numpy()
                    ac = r_all.boxes.conf.cpu().numpy()
                    order = np.argsort(-ac)
                    ious = np.array([iou_xywh(ab[j], gt[i]) for j in order])
                    for t in np.arange(0.5, 1.0, 0.05):
                        lst = dets.setdefault(round(float(t), 2), [])
                        hit = False
                        for c, v in zip(ac[order], ious >= t):
                            is_tp = int(v and not hit)
                            hit = hit or bool(v)
                            lst.append((float(c), is_tp))
                else:
                    for t in np.arange(0.5, 1.0, 0.05):
                        dets.setdefault(round(float(t), 2), [])
            else:
                fp += len(boxes)
            done += 1
            pct = 100 * done // total
            if pct != last_pct and (pct % 10 == 0 or done == total):
                print(f"    [{tag}] {done}/{total} ({pct}%)", flush=True)
                last_pct = pct
        cap.release()
    p = tp / (tp + fp + 1e-9)
    r = tp / (tp + fn + 1e-9)
    f1 = 2 * p * r / (p + r + 1e-9)
    map50 = ap_voc(dets.get(0.5, []), n_gt)  # один клас -> AP = mAP@50
    aps = [ap_voc(dets.get(round(float(t), 2), []), n_gt) for t in np.arange(0.5, 1.0, 0.05)]
    map5095 = float(np.mean(aps)) if aps else 0.0
    ms = infer_t / max(frames, 1) * 1000
    return dict(gt=n_gt, tp=tp, fp=fp, fn=fn, frames=frames, p=p, r=r, f1=f1,
                ap=map50, map50=map50, map5095=map5095,
                ms=ms, fps=1000 / ms if ms else 0)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ai", default="all", help="номери моделей 1-7 через кому, або all")
    ap.add_argument("--val", default="val", help="папка валідації")
    ap.add_argument("--modality", default="visible", choices=["visible", "infrared", "both"])
    ap.add_argument("--conf", type=float, required=True)
    ap.add_argument("--iou", type=float, default=0.45)
    ap.add_argument("--imgsz", type=int, default=640)
    ap.add_argument("--stride", type=int, default=10, help="кожен N-й кадр (1 = всі кадри)")
    ap.add_argument("--max-seqs", type=int, default=0, help="обмежити к-сть сцен (0 = всі)")
    ap.add_argument("--device", default=None)
    ap.add_argument("--out", default="runs/eval_detect.xlsx")
    args = ap.parse_args()

    device = args.device or ("cuda:0" if torch.cuda.is_available() else "cpu")
    mods = [args.modality] if args.modality != "both" else ["visible", "infrared"]
    seq_dirs = sorted(p for p in (ROOT / args.val).iterdir() if p.is_dir())
    if args.max_seqs:
        seq_dirs = seq_dirs[:args.max_seqs]

    out = ROOT / args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")
        print(f"  [info] пишу в Excel: {out.name}")

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    def save_row(row):
        """Дописує рядок у xlsx (створює з заголовками при першому запуску)."""
        if out.exists():
            wb = load_workbook(out)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "detect"
            ws.append([h for _, h in COLUMNS])
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            ws.sheet_properties.pageSetUpPr = None
        ws.append([row.get(k) for k, _ in COLUMNS])
        # ширина колонок під вміст
        for i, (k, h) in enumerate(COLUMNS, 1):
            w = max(len(h), max((len(str(ws.cell(r, i).value or "")) for r in range(1, ws.max_row + 1)), default=0))
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(w + 2, 45)
        try:
            wb.save(out)
            return out
        except PermissionError:
            import datetime
            alt = out.with_name(f"{out.stem}_{datetime.datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
            print(f"  [warn] {out.name} зайнятий (відкритий в Excel?) — пишу в {alt.name}")
            wb.save(alt)
            return alt

    total_new = 0
    for num in parse_ai(args.ai):
        mp = ROOT / "ai" / MODELS[num]
        if not mp.exists():
            print(f"  [skip] модель {num}: немає файлу {mp}")
            continue
        seqs = []
        n_skip = 0
        for s in seq_dirs:
            for m in mods:
                g = load_gt(s, m)
                if g:
                    seqs.append((s.name, m, *g))
                else:
                    n_skip += 1
        if n_skip:
            print(f"  [skip] пар mp4+json немає: {n_skip}")
        if not seqs:
            print(f"  [skip] модель {num}: нема сцен для тестування")
            continue
        print(f"\n=== модель {num}: {MODELS[num]} | conf={args.conf} iou={args.iou} imgsz={args.imgsz} ===")
        print(f"  відео: {len(seqs)}")
        model = YOLO(str(mp))
        params, size_mb, _ = model_info(mp, model)
        tag = f"AI{num} c={args.conf} iou={args.iou} sz={args.imgsz}"
        m = eval_combo(model, [(v, e, g) for _, _, v, e, g in seqs],
                       args.conf, args.iou, args.imgsz, device, args.stride, tag)
        _, _, vram = model_info(mp, model)  # пік VRAM за прогін (на CUDA; на CPU 0)
        row = dict(model=num, weights=MODELS[num], modality=args.modality,
                   conf=args.conf, iou=args.iou, imgsz=args.imgsz,
                   seqs=len(seqs), params=round(params, 2), size=round(size_mb, 1),
                   device=device, vram=round(vram, 2), **m)
        saved_to = save_row(row)
        total_new += 1
        print(f"  P={m['p']:.3f} R={m['r']:.3f} F1={m['f1']:.3f} "
              f"mAP@50={m['map50']:.3f} mAP@50:95={m['map5095']:.3f} "
              f"TP={m['tp']} FP={m['fp']} FN={m['fn']} GT={m['gt']} "
              f"{params:.1f}M {size_mb:.1f}MB [{device}] {m['fps']:.1f}fps")

    print(f"\n[saved] {saved_to if total_new else out} (+{total_new} рядків)")


if __name__ == "__main__":
    main()
